#!/usr/bin/env python3
"""Generate test/fixtures/parser/die-layout.xlsx (基板布局 Excel).

One sheet per product, in the format parse_die_layout_xls expects:
  - row 1: blank corner cell, then X grid coordinates
  - column 1 (from row 2): Y grid coordinates
  - cells: '1' where the layout has a die, 'S' at alignment markers,
    empty where there is no die

The grids are derived from the parsed CP1 layers of the fixture lots
(test/fixtures/parsed/<lot>/cp1.json), so the alignment markers match the
wafer maps exactly and layout-adoption succeeds end to end.

Rerun after regenerating parsed fixtures:
    python3 test/fixtures/tools/generate_die_layout.py
"""
import json
from pathlib import Path

from openpyxl import Workbook

FIXTURES = Path(__file__).resolve().parents[1]

# sheet name (product id) -> parsed cp1 layer to derive the grid from
SHEETS = {
    "S1M032120B": FIXTURES / "parsed" / "B003332" / "cp1.json",
    "S1M040120B": FIXTURES / "parsed" / "B003990" / "cp1.json",
}

MARKERS = {"S", "*"}


def cell_value(die):
    bin_value = die["bin"]
    special = bin_value.get("special")
    if special in MARKERS:
        return "S"
    return "1"


def main() -> None:
    wb = Workbook()
    wb.remove(wb.active)

    for product, source in SHEETS.items():
        dies = json.loads(source.read_text())["map"]["dies"]
        xs = sorted({d["x"] for d in dies})
        ys = sorted({d["y"] for d in dies})
        by_pos = {(d["x"], d["y"]): d for d in dies}

        ws = wb.create_sheet(title=product)
        for col, x in enumerate(xs, start=2):
            ws.cell(row=1, column=col, value=x)
        for row, y in enumerate(ys, start=2):
            ws.cell(row=row, column=1, value=y)
            for col, x in enumerate(xs, start=2):
                die = by_pos.get((x, y))
                if die is not None:
                    ws.cell(row=row, column=col, value=cell_value(die))

        markers = sum(
            1 for d in dies if d["bin"].get("special") in MARKERS
        )
        print(f"sheet {product}: {len(dies)} dies, {markers} markers")

    out = FIXTURES / "parser" / "die-layout.xlsx"
    wb.save(out)
    print(f"wrote {out}")


if __name__ == "__main__":
    main()
